from io import BytesIO
import json
import sys
import types


def test_normalize_and_offline_story_analysis():
    from storyforge.parser import normalize_text
    from storyforge.ai.heuristic_provider import HeuristicProvider

    normalized, paragraphs = normalize_text("Maya enters the Station.\n\nA secret threatens her.")
    assert normalized == "Maya enters the Station.\nA secret threatens her."
    assert len(paragraphs) == 2

    provider = HeuristicProvider()
    bible = provider.analyze_story(normalized, "Test Story")
    assert bible["title"] == "Test Story"
    assert bible["major_events"]
    assert bible["characters"]
    planned = provider.plan_scenes(bible, normalized)
    assert len(planned) >= 3


def test_document_parsers_all_formats():
    from docx import Document
    from reportlab.pdfgen import canvas
    from storyforge.parser import parse_uploaded_file

    text, _ = parse_uploaded_file("story.txt", b"A simple story.")
    assert "simple story" in text

    doc = Document()
    doc.add_paragraph("DOCX story")
    bio = BytesIO()
    doc.save(bio)
    text, _ = parse_uploaded_file("story.docx", bio.getvalue())
    assert "DOCX story" in text

    pdf = BytesIO()
    c = canvas.Canvas(pdf)
    c.drawString(72, 720, "PDF story")
    c.save()
    text, _ = parse_uploaded_file("story.pdf", pdf.getvalue())
    assert "PDF story" in text


def test_local_json_extraction():
    from storyforge.ai.local_provider import LocalLLMProvider

    raw = "Here is the JSON:\n```json\n{\"title\":\"Demo\",\"major_events\":[]}\n```"
    value = LocalLLMProvider._extract_json(raw)
    assert value["title"] == "Demo"


def test_local_provider_with_fake_llama(tmp_path, monkeypatch):
    class FakeLlama:
        def __init__(self, **kwargs):
            self.kwargs = kwargs
            self.calls = []

        def create_chat_completion(self, **kwargs):
            self.calls.append(kwargs)
            user = kwargs["messages"][-1]["content"]
            if "object with a scenes array" in user:
                payload = {"scenes": [{
                    "scene_number": 1, "act": 1, "sequence": 1, "slugline": "INT. ROOM - DAY",
                    "interior_or_exterior": "INT.", "location": "ROOM", "time_of_day": "DAY",
                    "characters": [], "story_purpose": "Introduce the conflict.", "source_event_ids": ["event-1"],
                    "emotional_tone": "tense", "transitions": [], "props": [], "production_flags": [],
                    "estimated_duration": 2,
                }]}
            elif "complete generated screenplay scene object" in user or "complete revised screenplay scene object" in user:
                payload = {
                    "action": "A character enters.",
                    "dialogue": [],
                    "transitions": [],
                    "props": [],
                    "production_flags": [],
                    "emotional_tone": "tense",
                    "screenplay_text": "INT. ROOM - DAY\n\nA character enters.",
                }
            else:
                payload = {
                    "title": "Demo",
                    "genres": ["Drama"],
                    "summary": "A short story.",
                    "themes": ["change"],
                    "narrative_tone": "dramatic",
                    "characters": [],
                    "locations": ["ROOM"],
                    "timeline": [],
                    "major_events": [{"id": "event-1", "order": 1, "description": "A key event", "source_paragraph_ids": []}],
                    "open_plot_threads": [],
                }
            return {"choices": [{"message": {"content": json.dumps(payload)}}]}

    monkeypatch.setitem(sys.modules, "llama_cpp", types.SimpleNamespace(Llama=FakeLlama))
    model_path = tmp_path / "fake.gguf"
    model_path.write_bytes(b"fake")

    from storyforge.ai.local_provider import LocalLLMProvider

    provider = LocalLLMProvider(model_path)
    bible = provider.analyze_story("A short story.", "Demo")
    assert bible["title"] == "Demo"
    planned = provider.plan_scenes(bible, "A short story.")
    assert planned[0]["slugline"] == "INT. ROOM - DAY"
    scene = provider.generate_scene(bible, planned[0], "A short story.")
    assert "screenplay_text" in scene
    assert provider.model.calls[0]["response_format"] == {"type": "json_object"}


def test_factory_caches_provider(monkeypatch):
    from storyforge.ai import factory

    created = []

    class FakeProvider:
        is_real_ai = True

        def __init__(self, path=None):
            self.path = path
            created.append(path)

    factory.clear_provider_cache()
    monkeypatch.setattr(factory, "LocalLLMProvider", FakeProvider)
    a = factory.get_provider("C:/models/test.gguf")
    b = factory.get_provider("C:/models/test.gguf")
    assert a is b
    assert created == ["C:/models/test.gguf"]
    factory.clear_provider_cache()


def test_ollama_provider_uses_json_mode_and_local_endpoint(monkeypatch):
    import io
    from storyforge.ai.ollama_provider import OllamaProvider

    captured = {}

    class FakeResponse:
        def __enter__(self):
            return self
        def __exit__(self, *args):
            return False
        def read(self):
            return json.dumps({
                "message": {"content": '{"title":"Test","genres":[],"summary":"ok","themes":[],"narrative_tone":"calm","characters":[],"locations":[],"timeline":[],"major_events":[],"open_plot_threads":[]}'},
                "eval_count": 12,
                "total_duration": 1_000_000_000,
            }).encode()

    def fake_urlopen(request, timeout=0):
        captured["url"] = request.full_url
        captured["payload"] = json.loads(request.data.decode())
        return FakeResponse()

    monkeypatch.setattr("urllib.request.urlopen", fake_urlopen)
    provider = OllamaProvider("gemma3:4b", "http://127.0.0.1:11434")
    response = provider._chat("return json", max_tokens=100)
    assert json.loads(response)["title"] == "Test"
    assert captured["url"] == "http://127.0.0.1:11434/api/chat"
    assert captured["payload"]["model"] == "gemma3:4b"
    assert captured["payload"]["format"] == "json"
    assert captured["payload"]["stream"] is False


def test_ollama_provider_recovers_when_scene_omits_screenplay_text(monkeypatch):
    responses = [
        {
            "message": {"content": json.dumps({
                "action": "", "dialogue": [], "transitions": [], "props": [],
                "production_flags": [], "emotional_tone": "tense"
            })},
            "eval_count": 100,
            "total_duration": 1_000_000_000,
        },
        {
            "message": {"content": json.dumps({
                "screenplay_text": (
                    "INT. OLD STATION - NIGHT\n\nMira enters the abandoned station. "
                    "Her flashlight catches a fresh footprint.\n\nMIRA\nWho's there?"
                )
            })},
            "eval_count": 100,
            "total_duration": 1_000_000_000,
        },
    ]
    requests = []

    class FakeResponse:
        def __init__(self, payload):
            self.payload = payload
        def __enter__(self):
            return self
        def __exit__(self, *args):
            return False
        def read(self):
            return json.dumps(self.payload).encode()

    def fake_urlopen(request, timeout=0):
        requests.append(json.loads(request.data.decode()))
        return FakeResponse(responses.pop(0))

    monkeypatch.setattr("urllib.request.urlopen", fake_urlopen)
    from storyforge.ai.ollama_provider import OllamaProvider

    provider = OllamaProvider("gemma3:4b", "http://127.0.0.1:11434")
    scene = provider.generate_scene(
        {"title": "Station Signal", "characters": [{"name": "Mira"}]},
        {
            "slugline": "INT. OLD STATION - NIGHT", "location": "OLD STATION",
            "story_purpose": "Mira discovers a footprint.", "characters": ["Mira"],
            "estimated_duration": 2,
        },
        "Mira enters an abandoned station and finds a footprint.",
    )
    assert "screenplay_text" in scene
    assert "Mira enters the abandoned station" in scene["screenplay_text"]
    assert len(requests) == 2
    assert "Write the complete screenplay text for this ONE scene" in requests[1]["messages"][1]["content"]


def test_story_bible_normalizes_dict_locations_and_list_fields():
    from storyforge.ai.local_provider import LocalLLMProvider
    from storyforge.text_utils import text_items, joined_text

    bible = LocalLLMProvider._normalize_bible({
        "title": "Demo",
        "genres": [{"name": "Mystery"}],
        "themes": [{"theme": "Trust"}],
        "locations": [{"name": "Old Station", "description": "An abandoned station"}],
        "open_plot_threads": [{"question": "Who sent the letter?"}],
        "characters": [{"name": "Mira"}],
        "major_events": [{"description": "Mira finds a letter"}],
        "timeline": [{"description": "Before dawn"}],
    }, "Demo")

    assert bible["locations"] == ["Old Station"]
    assert bible["genres"] == ["Mystery"]
    assert bible["themes"] == ["Trust"]
    assert bible["open_plot_threads"] == ["Who sent the letter?"]
    assert "Old Station" in "\n".join(text_items([{"name": "Old Station"}], kind="location"))
    assert joined_text([{"name": "Mystery"}], kind="genre") == "Mystery"



def test_database_serializes_list_values_before_sqlite_binding(tmp_path, monkeypatch):
    from storyforge import db

    monkeypatch.setattr(db, "DB_PATH", tmp_path / "storyforge.db")
    project = db.create_project("DB safety test")
    scene = {
        "scene_number": 1, "act": 1, "sequence": 1,
        "slugline": "INT. ROOM - DAY", "interior_or_exterior": "INT.",
        "location": "ROOM", "time_of_day": "DAY",
        "characters": [], "story_purpose": "Introduce story",
        "source_event_ids": [], "action": "A character enters.",
        "dialogue": [], "emotional_tone": "neutral", "transitions": [],
        "props": [], "production_flags": [], "estimated_duration": 2,
        "screenplay_text": "INT. ROOM - DAY\n\nA character enters.",
    }
    saved = db.replace_scenes(project["id"], [scene])[0]
    # Simulate malformed LLM output: action should be text but arrived as a list.
    updated = db.update_scene(saved["id"], {"action": ["A character enters.", "A phone rings."]})
    assert updated["action"] == '["A character enters.", "A phone rings."]'


def test_scene_text_is_composed_when_model_omits_screenplay_text():
    from storyforge.ai.local_provider import LocalLLMProvider

    plan = {
        "slugline": "INT. OLD STATION - NIGHT",
        "story_purpose": "Mira discovers who sent the warning.",
        "props": [],
        "production_flags": [],
    }
    result = {
        "action": "Mira pushes open the station door. Dust swirls in the flashlight beam.",
        "dialogue": [{"character": "MIRA", "text": "Who's there?"}],
    }
    text = LocalLLMProvider._compose_screenplay_text(result, plan, require_generated_parts=True)
    assert "INT. OLD STATION - NIGHT" in text
    assert "Dust swirls" in text
    assert "MIRA" in text
    assert "Who's there?" in text


def test_local_provider_retries_screenplay_text_only_when_missing(tmp_path, monkeypatch):
    class FakeLlama:
        def __init__(self, **kwargs):
            self.calls = []

        def create_chat_completion(self, **kwargs):
            self.calls.append(kwargs)
            prompt = kwargs["messages"][-1]["content"]
            if "Write the complete screenplay text for this ONE scene" in prompt:
                payload = {
                    "screenplay_text": (
                        "INT. OLD STATION - NIGHT\n\nMira enters the abandoned station. "
                        "Her flashlight catches a fresh footprint.\n\nMIRA\nWho's there?"
                    )
                }
            else:
                payload = {
                    "action": "",
                    "dialogue": [],
                    "transitions": [],
                    "props": [],
                    "production_flags": [],
                    "emotional_tone": "tense",
                }
            return {"choices": [{"message": {"content": json.dumps(payload)}}]}

    monkeypatch.setitem(sys.modules, "llama_cpp", types.SimpleNamespace(Llama=FakeLlama))
    model_path = tmp_path / "fake.gguf"
    model_path.write_bytes(b"fake")

    from storyforge.ai.local_provider import LocalLLMProvider

    provider = LocalLLMProvider(model_path)
    scene_plan = {
        "slugline": "INT. OLD STATION - NIGHT",
        "location": "OLD STATION",
        "story_purpose": "Mira discovers a footprint.",
        "characters": ["Mira"],
        "estimated_duration": 2,
    }
    generated = provider.generate_scene(
        {"title": "Station Signal", "characters": [{"name": "Mira"}]},
        scene_plan,
        "Mira enters an abandoned station and finds a footprint.",
    )
    assert "screenplay_text" in generated
    assert "Mira enters the abandoned station" in generated["screenplay_text"]
    assert len(provider.model.calls) == 2


def test_scene_generation_uses_linked_source_paragraphs_to_reduce_prompt_context():
    from storyforge.services.scenes import _scene_source_context

    paragraphs = [
        {"id": "p-1", "position": 1, "text": "Opening paragraph establishes the town. " + ("Opening detail. " * 50)},
        {"id": "p-2", "position": 2, "text": "Mira arrives at the station. " + ("Arrival detail. " * 50)},
        {"id": "p-3", "position": 3, "text": "Mira finds the brass key at the station. " + ("Relevant discovery detail. " * 50)},
        {"id": "p-4", "position": 4, "text": "She realizes the key belongs to her brother. " + ("Transition detail. " * 50)},
        {"id": "p-5", "position": 5, "text": "Mira leaves the station. " + ("Departure detail. " * 50)},
        {"id": "p-6", "position": 6, "text": "Much later, the mayor leaves the city. " + ("Unrelated ending detail. " * 50)},
    ]
    bible = {
        "major_events": [{"id": "event-2", "source_paragraph_ids": ["p-3"], "description": "Mira finds the key."}]
    }
    scene = {"source_event_ids": ["event-2"]}
    document = {
        "normalized_text": "\n".join(p["text"] for p in paragraphs),
        "source_paragraphs": paragraphs,
    }
    context = _scene_source_context(bible, scene, document)
    assert "RELEVANT SOURCE EXCERPTS" in context
    assert all(f"[p-{i}]" in context for i in range(1, 6))
    assert "[p-6]" not in context


def test_quality_findings_migrate_old_database_and_keep_evidence(tmp_path, monkeypatch):
    import sqlite3
    from storyforge import db

    db_path = tmp_path / "storyforge.db"
    old = sqlite3.connect(db_path)
    old.execute(
        "CREATE TABLE analysis_issues (id TEXT PRIMARY KEY, project_id TEXT NOT NULL, scene_id TEXT, "
        "check_type TEXT NOT NULL, severity TEXT NOT NULL, message TEXT NOT NULL, created_at TEXT NOT NULL)"
    )
    old.commit()
    old.close()
    monkeypatch.setattr(db, "DB_PATH", db_path)

    project = db.create_project("Issue detail migration")
    saved = db.replace_quality_issues(project["id"], [{
        "check_type": "timeline",
        "severity": "warning",
        "message": "A timeline detail may conflict.",
        "evidence": "Scene 2 is set before an event that Scene 1 says has already happened.",
        "suggested_action": "Recheck the order of these scenes.",
    }])
    assert saved[0]["evidence"].startswith("Scene 2")
    assert saved[0]["suggested_action"] == "Recheck the order of these scenes."


def test_quality_prompt_requires_actionable_review_findings():
    from storyforge.ai.prompts import quality_prompt

    prompt = quality_prompt(
        {"title": "Demo", "summary": "A mystery"},
        [{"scene_number": 1, "slugline": "INT. ROOM - NIGHT", "screenplay_text": "A scene." * 1000}],
        "Mira finds a key.",
    )
    assert '"suggested_action"' in prompt
    assert "never return generic placeholders" in prompt
    assert '"screenplay_text"' in prompt


def test_fast_quality_check_skips_the_expensive_ai_review(monkeypatch):
    from storyforge.services import quality

    scene = {
        "id": "scene-1", "scene_number": 1, "slugline": "INT. ROOM - DAY",
        "location": "ROOM", "characters": [], "source_event_ids": [],
        "story_purpose": "Introduce the room.", "screenplay_text": "INT. ROOM - DAY\n\nA character enters.",
        "action": "A character enters and looks around the room.", "dialogue": [], "production_flags": [],
    }
    class Provider:
        is_real_ai = True
        def quality_review(self, *args):
            raise AssertionError("AI review must not run in fast mode")

    monkeypatch.setattr(quality, "get_story_bible", lambda _project: {
        "characters": [], "locations": [], "major_events": [],
    })
    monkeypatch.setattr(quality.db, "list_scenes", lambda _project: [scene])
    monkeypatch.setattr(quality, "latest", lambda _project: {"normalized_text": "A story."})
    monkeypatch.setattr(quality, "get_provider", lambda _model: Provider())
    monkeypatch.setattr(quality.db, "replace_quality_issues", lambda _project, issues: issues)

    assert quality.run("project-1", "ollama:gemma3:4b", include_ai_review=False) == []


def test_ai_quality_findings_keep_specific_message_evidence_and_action(monkeypatch):
    from storyforge.services import quality

    scene = {
        "id": "scene-1", "scene_number": 1, "slugline": "INT. ROOM - DAY",
        "location": "ROOM", "characters": [], "source_event_ids": [],
        "story_purpose": "Introduce the room.", "screenplay_text": "INT. ROOM - DAY\n\nA character enters.",
        "action": "A character enters and looks around the room.", "dialogue": [], "production_flags": [],
    }
    class Provider:
        is_real_ai = True
        def quality_review(self, *args):
            return [{
                "scene_number": 1,
                "check_type": "timeline",
                "severity": "warning",
                "message": "Scene 1 occurs after the same character is already described as leaving.",
                "evidence": "The source says the character leaves before this scene begins.",
                "suggested_action": "Move the scene earlier or revise its time reference.",
            }]

    monkeypatch.setattr(quality, "get_story_bible", lambda _project: {
        "characters": [], "locations": [], "major_events": [],
    })
    monkeypatch.setattr(quality.db, "list_scenes", lambda _project: [scene])
    monkeypatch.setattr(quality, "latest", lambda _project: {"normalized_text": "A story."})
    monkeypatch.setattr(quality, "get_provider", lambda _model: Provider())
    monkeypatch.setattr(quality.db, "replace_quality_issues", lambda _project, issues: issues)

    findings = quality.run("project-1", "ollama:gemma3:4b", include_ai_review=True)
    ai_finding = next(item for item in findings if item["check_type"] == "timeline")
    assert "already described as leaving" in ai_finding["message"]
    assert "source says" in ai_finding["evidence"]
    assert ai_finding["suggested_action"].startswith("Move the scene earlier")
